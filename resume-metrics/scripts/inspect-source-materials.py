from __future__ import annotations

import hashlib
import json
import sys
from pathlib import Path

import pdfplumber
from pptx import Presentation


SUPPORTED_UPLOADS = {".pdf", ".pptx", ".txt", ".md", ".mp3", ".mp4"}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def slide_text(slide) -> str:
    values: list[str] = []
    for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text:
            value = " ".join(shape.text.split())
            if value:
                values.append(value)
    return "\n".join(values)


def inspect_pptx(path: Path) -> dict:
    deck = Presentation(path)
    slides = []
    for index, slide in enumerate(deck.slides, start=1):
        text = slide_text(slide)
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.text:
            title = " ".join(slide.shapes.title.text.split())
        slides.append({
            "slide": index,
            "title": title,
            "text_chars": len(text),
            "text_preview": text[:500],
        })
    return {
        "kind": "pptx",
        "slide_count": len(slides),
        "nonempty_slides": sum(1 for item in slides if item["text_chars"] > 0),
        "text_chars": sum(item["text_chars"] for item in slides),
        "slides": slides,
    }


def inspect_pdf(path: Path) -> dict:
    pages = []
    with pdfplumber.open(path) as document:
        for index, page in enumerate(document.pages, start=1):
            text = page.extract_text() or ""
            normalized = "\n".join(line.strip() for line in text.splitlines() if line.strip())
            pages.append({
                "page": index,
                "width": round(float(page.width), 2),
                "height": round(float(page.height), 2),
                "text_chars": len(normalized),
                "text_preview": normalized[:700],
            })
    return {
        "kind": "pdf",
        "page_count": len(pages),
        "nonempty_pages": sum(1 for item in pages if item["text_chars"] > 0),
        "text_chars": sum(item["text_chars"] for item in pages),
        "pages": pages,
    }


def inspect_file(path: Path) -> dict:
    item = {
        "name": path.name,
        "path": str(path.resolve()),
        "extension": path.suffix.lower(),
        "size_bytes": path.stat().st_size,
        "sha256": sha256(path),
        "accepted_by_current_upload_policy": path.suffix.lower() in SUPPORTED_UPLOADS,
    }
    if path.suffix.lower() == ".pptx":
        item.update(inspect_pptx(path))
    elif path.suffix.lower() == ".pdf":
        item.update(inspect_pdf(path))
    else:
        item["kind"] = "audio" if path.suffix.lower() in {".m4a", ".mp3"} else "other"
    return item


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("usage: inspect-source-materials.py SOURCE_DIR OUTPUT_JSON")
    source = Path(sys.argv[1])
    output = Path(sys.argv[2])
    files = [inspect_file(path) for path in sorted(source.iterdir()) if path.is_file()]
    summary = {
        "file_count": len(files),
        "pptx_count": sum(1 for item in files if item["kind"] == "pptx"),
        "pdf_count": sum(1 for item in files if item["kind"] == "pdf"),
        "audio_count": sum(1 for item in files if item["kind"] == "audio"),
        "total_slides": sum(item.get("slide_count", 0) for item in files),
        "total_pdf_pages": sum(item.get("page_count", 0) for item in files),
        "total_extractable_text_chars": sum(item.get("text_chars", 0) for item in files),
        "unsupported_files": [item["name"] for item in files if not item["accepted_by_current_upload_policy"]],
    }
    result = {"summary": summary, "files": files}
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
